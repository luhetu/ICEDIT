"""Build the dated, evidence-qualified tutor progress presentation."""
from pathlib import Path
from io import BytesIO
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "reports"
INK, GREEN, MUTED = "20282C", "16776B", "647176"
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)


def text(slide, x, y, w, h, value, size=22, color=INK, bold=False):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    for i, line in enumerate(value.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.font.name = "Calibri"
        p.font.size = Pt(size)
        p.font.bold = bold
        p.font.fill.solid()
        p.font.fill.fore_color.rgb = RGBColor.from_string(color)
        p.space_after = Pt(14)
    return box


def slide(title, subtitle):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    text(s, .55, .28, 12, .35, "ICEDIT / RESEARCH PROGRESS / 16 SEP 2026", 12, GREEN, True)
    text(s, .55, .87, 12.2, .85, title, 32, bold=True)
    text(s, .55, 1.77, 12, .7, subtitle, 18, MUTED)
    text(s, .55, 7.04, 11.6, .25, "CVPR-oriented research | Implemented components and hypotheses are distinguished", 11, MUTED)
    text(s, 12.15, 7.02, .6, .3, str(len(prs.slides)), 12, GREEN)
    return s


def panel(s, x, y, w, h, title, body, fill="EDF4F2"):
    shape = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor.from_string(fill)
    shape.line.fill.background()
    text(s, x+.2, y+.15, w-.4, .6, title, 22, GREEN, True)
    text(s, x+.2, y+.88, w-.4, h-.95, body, 18)


s = slide("Remove the Target. Preserve the Others.", "Research focus: selective object removal when target and protected content interact.")
panel(s, .6, 2.65, 5.9, 3.6, "The practical problem", "Remove the specified instance and its effects.\nPreserve nearby faces, text, structure and other instances.\nA plausible image can still be an incorrect edit.")
panel(s, 6.8, 2.65, 5.9, 3.6, "Our current stage", "Editing infrastructure is implemented.\nStrong-baseline validation is in progress.\nA new method contribution remains to be demonstrated.", "F1F2F5")

s = slide("What We Have Built", "Implemented pipeline; active denoising intervention remains a proposal.")
labels = ["Instruction", "Edit contract", "Spatial controls", "ICEdit / Fill", "Raw + composite"]
for i, label in enumerate(labels):
    x = .6 + i*2.57
    panel(s, x, 2.8, 2.25, 1.35, label, "")
    if i < 4:
        arrow = s.shapes.add_shape(MSO_SHAPE.CHEVRON, Inches(x+2.29), Inches(3.25), Inches(.22), Inches(.3))
        arrow.fill.solid(); arrow.fill.fore_color.rgb = RGBColor.from_string(GREEN)
        arrow.line.fill.background()
text(s, .75, 4.6, 11.8, 1.6, "Contract: target, operation and preservation requirements.\nControls: source-only masks and protected-region subtraction.\nFinal compositing preserves pixels outside the mask; it does not prove better generation.", 23)

s = slide("Progress and Evidence", "Verified September 16: one scene, five conditions, three seeds; not a broad benchmark.")
panel(s, .6, 2.6, 5.9, 3.9, "Completed ICEdit pilot", "15/15 generations completed; 30 infrastructure tests passed.\nOutside-mask raw L1: full 0.05176; masked 0.02246.\nAbout 57% less pixel change, but signs still remain or regenerate.")
panel(s, 6.8, 2.6, 5.9, 3.9, "Interpretation and status", "Pixel preservation does not imply successful removal or preserved identity.\nNative Fill removes rooftop signs but alters people.\nObjectClear imports successfully; GPU runs are still pending.", "F1F2F5")

s = slide("Where a Contribution Must Go Further", "A contract, verifier or final composite alone is not the proposed contribution.")
text(s, .75, 2.65, 11.8, 2.45, "ObjectClear: target/effect localization and background fusion.\nConsistency Critic: reference-guided consistency correction.\nAdaEraser: adaptive attention suppression for removal.", 25)
text(s, .75, 5.0, 11.8, .85, "Hypothesis: explicit instance ownership can reduce collateral damage when effects or boundaries overlap.", 25, GREEN, True)
text(s, .75, 6.18, 11.8, .55, "Sources: arxiv.org/abs/2505.22636v2 | arxiv.org/abs/2511.20614 | arxiv.org/abs/2605.15921", 12, MUTED)

s = slide("ObjectClear: What the Paper Does", "Precise Object and Effect Removal with Adaptive Target-Aware Attention | CVPR 2026")
panel(s, .6, 2.6, 5.9, 3.5, "Backbone and learning", "SDXL-Inpainting with a CLIP visual object encoder.\nInput: source image and target-object mask.\nOBER supplies paired scenes and object-effect supervision.")
panel(s, 6.8, 2.6, 5.9, 3.5, "Three mechanisms", "ATA: learn where the target and its effects are.\nSVDS: vary denoising strength across regions.\nAGF: use attention to fuse generated content with source details.", "F1F2F5")
text(s, .75, 6.3, 11.8, .45, "Zhao et al., arxiv.org/html/2505.22636v2 | Paper mechanisms, not our contributions", 14, MUTED)
s.notes_slide.notes_text_frame.text = "这一页介绍已有工作。ObjectClear 基于 SDXL-Inpainting，输入原图和目标 mask。它通过目标及影响的监督学习定位，再结合空间变化去噪和注意力融合保持背景。这些机制属于论文作者，不能作为我们的创新。"

s = slide("ObjectClear: Simplified Method Diagram", "Conceptual redraw based on the paper; training supervision is distinct from inference input.")
stages = [("Source + mask", "Select the target"), ("SDXL + ATA", "Remove object/effects"), ("AGF", "Fuse source details")]
for i, (heading, body) in enumerate(stages):
    x = .7 + i * 4.25
    panel(s, x, 2.75, 3.7, 1.8, heading, body)
    if i < 2:
        arrow = s.shapes.add_shape(MSO_SHAPE.CHEVRON, Inches(x+3.82), Inches(3.35), Inches(.3), Inches(.35))
        arrow.fill.solid(); arrow.fill.fore_color.rgb = RGBColor.from_string(GREEN)
        arrow.line.fill.background()
text(s, .8, 4.95, 11.6, 1.2, "SVDS acts during denoising; AGF uses the predicted attention map.\nOBER object-effect masks supervise training.", 23)
text(s, .8, 6.35, 11.6, .4, "Source: arxiv.org/html/2505.22636v2, Sec. 3.2 | Simplified, not an exact computation graph", 13, MUTED)
s.notes_slide.notes_text_frame.text = "流程图是我们依据论文重画的概念图。注意区分：推理输入是原图和目标 mask；物体加影响的标注用于训练监督。不能把它讲成推理时已经知道正确的影响范围。"

s = slide("ObjectClear: Our Reproduction and Research Gap", "Status checked September 16: submitted runs are pending; no reproduced quality claim yet.")
panel(s, .6, 2.6, 5.9, 3.85, "Our reproduction work", "Official code and FP16 weights prepared; pipeline import passes.\n1044511: official sample, AGF on/off.\n1044514: same ICEdit source/mask, 3 seeds x 2 modes; depends on smoke success.\nSave attention and pre/post-fusion outputs.")
panel(s, 6.8, 2.6, 5.9, 3.85, "What we still need to test", "Does a strong remover already solve our failures?\nDo simple protected masks suffice?\nCandidate gap: protecting other instances when deletion regions interact.\nAn explicit protection controller remains a hypothesis.", "F1F2F5")
text(s, .75, 6.57, 11.8, .32, "Code caveat: AGF toggle changes early latent blending AND final fusion. Use within-run pre/post pairs to isolate final fusion.", 12, MUTED)
s.notes_slide.notes_text_frame.text = "我们已完成下载、环境修复和导入验证，生成任务仍排队。首先确认强基线能否解决现有失败，再判断简单保护 mask 是否足够。候选方向是实例间误伤控制，但还没有证明新颖性或效果。AGF 开关还影响早期潜变量回填，因此不能直接把开关实验说成纯后处理消融；同一次运行的融合前后图才用于检查最终融合收益。"

s = slide("Candidate Mechanism: Test Before Building", "Proposed diagnostic, not a validated method or physical causal model.")
text(s, .75, 2.6, 11.8, 3.5, "1. Identify target, protected instances and uncertain regions.\n2. At the same denoising state, vary target and protected reference evidence separately.\n3. Test whether prediction responses distinguish target persistence from protected-content damage.\n4. Build a regional controller only if this signal beats simple spatial heuristics.", 24)
text(s, .75, 6.2, 11.8, .55, "Compare equal information and compute budgets. Report raw outputs before compositing.", 18, GREEN, True)

s = slide("Next Experiments and Decision Criteria", "Use existing cases first; no new benchmark is required for this diagnostic phase.")
panel(s, .6, 2.65, 5.9, 3.8, "Experimental sequence", "Completed: ICEdit, 5 conditions x 3 seeds.\nNext: validate ObjectClear and fusion effects.\nCompare simple protection controls on held-out cases.\nThen test the smallest intervention module.")
panel(s, 6.8, 2.65, 5.9, 3.8, "Continue only if...", "The failure persists on strong baselines.\nRemoval improves at comparable protection damage, or vice versa.\nGains survive raw-output and held-out evaluation.\nSimple masks and equal-cost resampling cannot explain the gain.", "F1F2F5")

OUT.mkdir(exist_ok=True)
path = OUT / "ICEdit_CVPR_Progress_2026-09-16.pptx"
buffer = BytesIO()
prs.save(buffer)
path.write_bytes(buffer.getvalue())
assert len(Presentation(path).slides) == 9
print(path)
